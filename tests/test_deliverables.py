import json
from pathlib import Path
import pandas as pd
from pptx import Presentation
from docx import Document
import pytest
import shutil
import nbformat


def test_real_deliverables_generation(completed_run):
    from src.reporting.deliverables import build_deliverables
    out,manifest=completed_run
    build_deliverables(out)
    dest=out/'deliverables'
    deck=Presentation(dest/'CRM_Final_Presentation.pptx')
    assert len(deck.slides)==12
    for slide in deck.slides:
        assert len(slide.notes_slide.notes_text_frame.text.split())>45
        for shape in slide.shapes:
            assert shape.left>=0 and shape.top>=0
            assert shape.left+shape.width<=deck.slide_width+100
            assert shape.top+shape.height<=deck.slide_height+100
    scored=pd.read_csv(out/'predictions/open_deal_predictions.csv')
    example=scored.loc[~scored.account_missing].iloc[0]
    output_slide=' '.join(shape.text for shape in deck.slides[8].shapes if shape.has_text_frame)
    assert f"{example['opportunity_id']} / {example['product']}" in output_slide
    assert 'bound method' not in output_slide
    text=' '.join(p.text for p in Document(dest/'CRM_Final_Report.docx').paragraphs)
    assert 'exploratory' in text.lower() and '1,589' in text
    paper=(dest/'CRM_Final_Report.md').read_text()
    for row in pd.read_csv(out/'metrics/test_metrics.csv').to_dict('records'):
        assert f"{row['roc_auc']:.4f}" in paper
    assert 'SMOKE_TEST_NOT_FINAL' in paper
    assert 'not causal' in paper.lower()
    assert (dest/'SPEAKER_SCRIPT_ESL.md').stat().st_size>3000


from tests.test_project_run import completed_run


@pytest.mark.parametrize('count',[0,2])
def test_documents_follow_actual_open_counts_including_empty(completed_run,tmp_path,count):
    from src.reporting.deliverables import build_deliverables
    original,_=completed_run
    out=tmp_path/'altered_run'; shutil.copytree(original,out)
    scores=pd.read_csv(out/'predictions/open_deal_predictions.csv').iloc[:count]
    scores.to_csv(out/'predictions/open_deal_predictions.csv',index=False)
    missing=int(scores.account_missing.sum())
    q=json.loads((out/'data/data_quality.json').read_text())
    q.update(scorable_open_rows=count,scorable_missing_account=missing,open_rows=count+3)
    q['raw_rows']['sales_pipeline']=q['rows']+count+3
    (out/'data/data_quality.json').write_text(json.dumps(q))
    manifest=json.loads((out/'run_manifest.json').read_text()); manifest['scored_open_rows']=count
    (out/'run_manifest.json').write_text(json.dumps(manifest))
    build_deliverables(out)
    summary=(out/'deliverables/RESULTS_SUMMARY.md').read_text()
    assert f'Open scores: {count:,}; missing-account warnings: {missing:,}' in summary
    paper=(out/'deliverables/CRM_Final_Report.md').read_text()
    assert '1,589' not in paper and '1,088' not in paper
    if count==0: assert 'No eligible Engaging opportunities' in paper
    deck=Presentation(out/'deliverables/CRM_Final_Presentation.pptx')
    texts=' '.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert f'{count:,} Engaging records scored; 3 Prospecting excluded' in texts
    notebook=nbformat.read(Path(__file__).resolve().parents[1]/'notebooks/CRM_Sales_Opportunities.ipynb',as_version=4)
    verification=next(cell.source for cell in notebook.cells if cell.cell_type=='code' and cell.source.startswith('import csv'))
    # Execute the real verification cell against a valid changed-size output.
    exec(verification,{'OUTPUT':out})
