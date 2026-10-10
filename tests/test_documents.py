"""The hand-maintained paper, slides, transcript and guide must cite the numbers of the saved run."""
import json
import re
import pandas as pd
import pytest
import os
from pathlib import Path
from src.outputs import DEFAULT_OUTPUT, ROOT

R=Path(os.environ.get('CRM_DOC_RUN',DEFAULT_OUTPUT))   # override only to check a copy of the run

pytestmark=pytest.mark.skipif(not (R/'agent/experiment.json').exists(),reason='full run not present')


def paper():
    from docx import Document
    d=Document(ROOT/'CRM_IEEE_Paper.docx')
    cells=[c.text for t in d.tables for r in t.rows for c in r.cells]
    return ' '.join([p.text for p in d.paragraphs]+cells)


def deck():
    from pptx import Presentation
    d=Presentation(ROOT/'CRM_Final.pptx'); texts=[]
    for s in d.slides:
        for sh in s.shapes:
            if sh.has_text_frame: texts.append(sh.text_frame.text)
            if sh.has_table: texts+=[c.text for r in sh.table.rows for c in r.cells]
    return ' '.join(texts),d


def read(name): return (ROOT/name).read_text(encoding='utf-8')


@pytest.fixture(scope='module')
def facts():
    m=json.loads((R/'run_manifest.json').read_text()); sel=m['selected_model']
    test=pd.read_csv(R/'metrics/test_metrics.csv').set_index('model')
    cmp_=pd.read_csv(R/'metrics/feature_set_comparison.csv')
    card=json.loads((R/'models/model_card.json').read_text())
    s=pd.read_csv(R/'agent/system_comparison.csv').set_index('system')
    e=json.loads((R/'agent/experiment.json').read_text())
    scen=pd.read_csv(R/'agent/scenario_results.csv').groupby('system').passed.sum()
    rep=pd.read_csv(R/'agent/replay_outcomes.csv').set_index(['system_label','group'])
    op=pd.read_csv(R/'predictions/open_deal_predictions.csv')
    unsupported=lambda r:int(r.unsupported_labels_out_of_support+r.unsupported_labels_borderline)
    raw_lr=cmp_.loc[(cmp_.feature_set=='raw')&(cmp_.model=='logistic_regression')].iloc[0]
    return {'feature_set':m['feature_set'],'selected':sel,'test_auc4':f"{test.loc[sel,'roc_auc']:.4f}",'test_auc3':f"{test.loc[sel,'roc_auc']:.3f}",
            'raw_lr_test':f'{raw_lr.test_roc_auc:.4f}','val_auc':f"{card['validation_roc_auc']:.3f}",
            'open':f'{len(op):,}','missing':f'{int(op.account_missing.sum()):,}',
            'unsup_A':f"{unsupported(s.loc['A']):,}",'unsup_B':f"{unsupported(s.loc['B']):,}",'unsup_C':unsupported(s.loc['C']),
            'recall_B':f"{s.loc['B','data_issue_recall']:.2f}",'manual_C':f"{s.loc['C','manual_steps_mean']:.2f}",
            'determinism':f"{e['determinism']['agent_identical_hashes']}/{e['determinism']['agent_cases']}",
            'replay':f"{e['tool_replay']['matched']}/{e['tool_replay']['calls']}",
            'scen_C':f"{int(scen['C'])}/10",'scen_B':f"{int(scen['B'])}/10",'scen_A':f"{int(scen['A'])}/10",
            'escalated_win':f"{rep.loc[('C_decision','ESCALATE_AT_RISK_REVIEW'),'observed_win_rate']:.1%}",
            'claims':f"{int(s.loc['C','claims_checked']):,}"}


def test_run_supports_the_story_the_documents_tell(facts):
    assert facts['feature_set']=='history' and facts['selected']=='logistic_regression'
    assert facts['unsup_C']==0


@pytest.mark.parametrize('name',['NOVELTY_AND_RESULTS.md','slide.md','docs/ESL_PRESENTATION_TRANSCRIPT.md','paper','deck'])
def test_core_numbers_appear_in_every_document(facts,name):
    text=paper() if name=='paper' else deck()[0] if name=='deck' else read(name)
    required=[facts['val_auc'],facts['open'],facts['missing'],facts['unsup_A'],facts['unsup_B']]
    required.append(facts['test_auc3'] if name=='deck' else facts['test_auc4'])
    if name in ('NOVELTY_AND_RESULTS.md','paper','deck'): required+=[facts['determinism'],facts['replay']]
    if name in ('NOVELTY_AND_RESULTS.md','paper'): required+=[facts['scen_C'],facts['claims'],facts['escalated_win'],facts['manual_C']]
    if name in ('NOVELTY_AND_RESULTS.md','slide.md'): required.append(facts['raw_lr_test'])
    missing=[v for v in required if v not in text]
    assert not missing,f'{name} lacks {missing}'


def test_guide_reference_numbers_and_commands_are_real(facts):
    guide=read('docs/EXCELLENT_PROJECT_DEMO_STEP_BY_STEP.md')
    for v in [facts['test_auc4'],facts['val_auc'],facts['open'],facts['missing'],facts['unsup_A'],facts['unsup_B'],
              facts['determinism'],facts['replay'],facts['escalated_win'],facts['scen_C'],facts['scen_B'],facts['scen_A']]:
        assert v in guide,v
    import importlib.util
    for module in set(re.findall(r'-m (src(?:\.\w+)+)',guide)):
        assert importlib.util.find_spec(module) is not None,module
    from src import agent_demo
    source=(ROOT/'src/agent_demo.py').read_text()
    for flag in set(re.findall(r'agent_demo[^\n`|]*?(--[a-z-]+)',guide)):
        assert f"'{flag}'" in source,flag
    for scenario in set(re.findall(r'--scenario ([a-z-]+)',guide)):
        assert scenario in agent_demo.SCENARIOS,scenario
    for path in set(re.findall(r'`(reports/crm/final/[^`*<]+?)`',guide)):
        assert (R/path.removeprefix('reports/crm/final/')).exists(),path


def test_deck_structure_and_notes():
    text,d=deck()
    assert len(d.slides)==24
    for i,slide in enumerate(d.slides,1):
        notes=slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ''
        assert len(notes.split())>=10,f'slide {i} has no speaker notes'
        for sh in slide.shapes:
            assert sh.left>=0 and sh.top>=0 and sh.left+sh.width<=d.slide_width+100 and sh.top+sh.height<=d.slide_height+100,(i,sh.name)
    assert 'python -m src.agent_demo' in text and 'First draft' not in paper()
