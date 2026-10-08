import json
from pptx import Presentation
from tests.test_experiments import comparison


def test_upgrade_artifacts_match_all_results(comparison):
    import pandas as pd
    from src.upgrade_deliverables import build_deliverables
    out,_=comparison
    build_deliverables(out)
    dest=out/'deliverables'; report=(dest/'CRM_Final_Report.md').read_text()
    for row in pd.read_csv(out/'ablation_comparison.csv').to_dict('records'):
        assert f"{row['test_roc_auc']:.4f}" in report
        assert f"{row['validation_roc_auc']:.4f}" in report
    assert 'frozen training archive' in report.lower()
    assert 'not causal' in report.lower()
    content=json.loads((dest/'slide_content.json').read_text())
    assert len(content)==12
    assert sum(s['seconds'] for s in content)==660
    for speaker in range(3):
        assert sum(s['seconds'] for s in content[speaker*4:speaker*4+4])==220
    assert all(len(s['note'].split())>=65 for s in content)
    deck=Presentation(dest/'CRM_Final_Presentation.pptx')
    assert len(deck.slides)==12
    assert len(deck.slides[5].shapes)>0
    all_text=' '.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert '[[' not in all_text
    assert json.loads((out/'demo_example.json').read_text())['opportunity']['opportunity_id'] in all_text
