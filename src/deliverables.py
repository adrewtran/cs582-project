"""Build editable course deliverables from the current run, never legacy results."""
from pathlib import Path
import argparse
import json
import pandas as pd
from docx import Document
from docx.shared import Inches as DInches,Pt as DPt
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

TEAM=['Hong Thai Phan','Nguyen Khanh An Tran','Hoang Thien Bao Bui']
REFERENCES=[
    '[1] Maven Analytics. CRM Sales Opportunities. Fictitious B2B computer-hardware company; four data tables. https://mavenanalytics.io/data-playground/crm-sales-opportunities (accessed 2026-10-06). Counts in this report are computed from the supplied CSV files, not website metadata.',
    '[2] J. Yan, M. Gong, C. Sun, J. Huang, and S. M. Chu. Sales Pipeline Win Propensity Prediction: A Regression Approach. IFIP/IEEE IM, 2015. https://arxiv.org/abs/1502.06229',
    '[3] A. Rezazadeh. A Generalized Flow for B2B Sales Predictive Modeling: An Azure Machine Learning Approach. Forecasting, 2020. https://doi.org/10.3390/forecast2030015 ; author manuscript: https://arxiv.org/abs/2002.01441',
    '[4] S. O. Arik and T. Pfister. TabNet: Attentive Interpretable Tabular Learning. AAAI 35(8), 6679–6687, 2021. https://doi.org/10.1609/aaai.v35i8.16826',
    '[5] scikit-learn. Probability calibration. https://scikit-learn.org/stable/modules/calibration.html (accessed 2026-10-06).',
    '[6] SHAP documentation. TreeExplainer. https://shap.readthedocs.io/en/latest/generated/shap.TreeExplainer.html (accessed 2026-10-06).',
    '[7] DreamQuark. pytorch-tabnet implementation and documentation. https://github.com/dreamquark-ai/tabnet (version 4.1.0 used).',
]


def pretty(name):
    return {'dummy_prior':'Dummy prior','logistic_regression':'Logistic Regression','random_forest':'Random Forest','mlp':'MLP','tabnet':'TabNet'}[name]


def paper_text(m,quality,table,validation,calibration,priority,history,example,shap):
    selected=table.set_index('model').loc[m['selected_model']]
    dummy=table.set_index('model').loc['dummy_prior']
    c=calibration.set_index('variant'); s=m['split']
    columns=['model','accuracy','precision','recall','f1','roc_auc','brier']
    results=table[columns].to_markdown(index=False,floatfmt='.4f')
    val=validation[['model','roc_auc','brier','threshold']].to_markdown(index=False,floatfmt='.4f')
    counts=table[['model','tn','fp','fn','tp','balanced_accuracy','macro_f1','average_precision']].to_markdown(index=False,floatfmt='.4f')
    settings='\n'.join(f"- {pretty(name)}: {json.dumps(values['settings'])}" for name,values in history.items())
    limitations='\n'.join('- '+item for item in m['limitations'])
    if example is None:
        example_text='No eligible Engaging opportunities were present in this run. The score CSV contains column headers and zero rows; no example probability is invented.'
    else:
        example_text=(f"A real exported example is opportunity {example.opportunity_id}, product {example['product']}: "
                      f"win probability {example.win_probability:.2%}, loss probability {example.loss_probability:.2%}. Its outcome is unknown.")
    return f'''# Predicting CRM Sales Opportunities Using Machine Learning

CS 582 • Group 2 • {', '.join(TEAM)}

Result-driven final paper draft for team review. Experiment mode: **{m['mode']}**. Generated from the run starting {m['started_utc']}. This document has not been submitted. Team members must verify the narrative, authorship/contributions and instructor formatting requirements before submission.

## Abstract

We implement a reproducible CRM classification and explanation pipeline using the supplied Maven Sales Opportunities data. The system estimates P(Won) and P(Lost) and reports the input factors behind each score. We compare a prior-only control, Logistic Regression, Random Forest, MLP and TabNet using identical dated partitions and training-only preprocessing. To reduce future-label leakage, a training or validation deal is eligible only when its outcome was available before the next period. The resulting split contains {s['train_rows']:,} training, {s['validation_rows']:,} validation and {s['test_rows']:,} test deals, with {s['purged_rows']:,} late-label records purged. {pretty(m['selected_model'])} is selected on validation ROC-AUC and reaches test ROC-AUC {selected.roc_auc:.4f}. These results do not establish practically useful discrimination. A deliberately invalid close-value control demonstrates how outcome leakage can create a misleading near-perfect score. We provide calibrated snapshot scores, auditable local sensitivities, Random Forest SHAP diagnostics and explicit input-quality warnings. The contribution is an applied, explainable evaluation workflow and an honest negative finding, not a new learning algorithm or proven sales uplift.

## 1. Problem and intended use

Sales teams need to understand which opportunities may close successfully and why a model gives a particular score. Our target is binary: Won = 1 and Lost = 0 among observed closed opportunities. We seek a prediction at the engagement stage using attributes available before closing. Open records have unknown outcomes and are never assigned artificial Lost labels. This dataset does not support a fixed-horizon claim such as winning within 30 days; we model eventual recorded outcome conditional on closure being observed.

The expected outputs are (1) complementary win/loss probabilities and (2) the input factors associated with a higher or lower model score. A predicted class and simple priority band are added for demonstration. They are not a validated business intervention. In particular, changing an explanatory factor is not guaranteed to change the real sales outcome.

## 2. Related work and scope of contribution

Yan et al. [2] frame win propensity as a component of sales pipeline management and evaluate a data-driven approach on enterprise B2B data. Their work motivates our probability output; we do not reproduce their proprietary data or claim a comparable business benefit.

Rezazadeh [3] presents a B2B workflow connecting probabilistic training with prediction and decision boundaries. This motivates separating model fitting, evaluation and scoring. Our implementation uses a small, free CPU workflow rather than reproducing its Azure infrastructure or monetary evaluation.

Arik and Pfister [4] introduce sequential attention for tabular learning in TabNet. We compare an existing TabNet implementation [7] with simpler learners, without claiming that advanced architecture must improve this CRM dataset. We do not implement its self-supervised pretraining.

Our course-project contribution combines availability-aware splitting, an explicit leakage control, probability diagnostics, local explanations and missing-input warnings in one regenerable pipeline. Explainable CRM scoring already exists in the literature. We therefore describe our novelty as the design and audit of this particular applied workflow, rather than a previously unknown algorithm.

## 3. Data and preparation

Maven describes the source as a fictitious B2B computer-hardware company [1]. The supplied files contain {quality['raw_rows']['sales_pipeline']:,} opportunities, {quality['raw_rows']['accounts']} accounts, {quality['raw_rows']['products']} products and {quality['raw_rows']['sales_teams']} sales-team records, plus a data dictionary. The opportunity table is left-joined to product, account and team tables with many-to-one checks. Keys, stage values, numerical fields, dates and row counts are validated. Source CSV files are unchanged and SHA-256 hashes are saved in the run manifest.

The loader normalizes GTXPro to GTX Pro ({quality['product_name_fixes']:,} pipeline rows) and fixes the documented sector spelling technolgy. There are {quality['rows']:,} labeled records: {quality['positives']:,} Won and {quality['negatives']:,} Lost. The {quality['open_rows']:,} open records include {quality['scorable_open_rows']:,} Engaging opportunities eligible for snapshot scoring and {quality['open_rows']-quality['scorable_open_rows']:,} Prospecting records, which are excluded. Among scored open records, {quality['scorable_missing_account']:,} lack an account; {quality['unmatched_account_in_labeled']:,} labeled records lack matched account attributes. Coverage differences limit the interpretation of open-deal scores.

Seven categorical predictors cover product, product series, account sector/location, regional office, manager and sales agent. Eleven numeric predictors cover product list price, account revenue and employees, revenue per employee, establishment year, account age, engagement calendar fields and subsidiary status. Revenue is expressed in millions in the source and is multiplied by one million when computing revenue per employee. We do not use account name as a predictor. close_value, close_date, deal_stage and opportunity_id are excluded from the feature matrix. close_date is read only for label availability; deal_stage defines the target; IDs preserve traceability.

Categoricals use constant missing-value imputation followed by one-hot encoding with unknown-category handling. Numeric values use training medians. Logistic Regression, MLP and TabNet also standardize numeric columns using training statistics. Missing open categories unseen during training are encoded as unknowns, so a missing-account warning remains necessary. EDA outcome comparisons use training rows; full-data missingness is an input-quality audit, not a model-selection signal.

![Input coverage](../figures/eda_missingness.png)

## 4. Evaluation design

The 60th and 80th percentiles of ordered closed engagement dates define period boundaries. All rows sharing a boundary date stay in the same period. The validation period starts {s['validation_start']}; the test period starts {s['test_start']}. A training row must engage and close before the validation boundary. A validation row must engage during its period and close before the test boundary. Test rows engage on or after the test boundary and have observed outcomes in the dataset. Counts are {s['train_rows']:,}/{s['validation_rows']:,}/{s['test_rows']:,}; {s['purged_rows']:,} late outcomes are recorded in split_manifest.csv rather than silently reassigned.

The win fractions are {s['train_win_rate']:.2%} in training, {s['validation_win_rate']:.2%} in validation and {s['test_win_rate']:.2%} in test. This is a single exploratory historical evaluation. Earlier project exploration already examined this dataset, and all boundaries are constructed retrospectively from the observed closed sample. The split is not an untouched prospective test and cannot remove closed-only selection bias or end-of-snapshot censoring.

All models receive exactly the same rows and feature roles. The model is chosen using validation ROC-AUC, with lower Brier score and then model name breaking ties. Class thresholds maximize validation macro-F1 over 0.10–0.90 in increments of 0.01; ties prefer the threshold closest to 0.50. No test statistic is used by the selection code. Neural early stopping also uses validation AUC, so the shared validation sample has several roles; a larger dataset should reserve a separate calibration/tuning sample or use nested temporal folds.

We report Accuracy, Precision, Recall, F1, ROC-AUC and confusion matrices as promised. Precision, recall and F1 treat Won as positive. Balanced accuracy, macro-F1, average precision, Brier and log loss provide additional context. Average precision is explicitly reported as AP, not mislabeled trapezoidal PR-AUC. F1 and accuracy can look favorable for a majority-Won model; the prior control is therefore essential.

## 5. Models and computation

Dummy prior outputs the training win fraction. Logistic Regression is an L2-regularized linear baseline. Random Forest uses bootstrap trees. MLP uses two ReLU hidden layers (64, 32), Adam, alpha 0.001 and learning rate 0.001. Its best validation checkpoint is retained. TabNet uses CPU attention steps with n_d = n_a = 8, n_steps = 3, learning rate 0.02, batch size 256 and virtual batch size 64. Its best validation checkpoint is also retained. We use seed 42 and limit numerical threads. These are fixed, modest comparison budgets, not an exhaustive hyperparameter search.

Actual settings for this run:

{settings}

The environment is Python {m['python']}, scikit-learn {m['dependencies']['scikit-learn']}, CPU PyTorch {m['dependencies']['torch']} and pytorch-tabnet {m['dependencies']['pytorch-tabnet']}. The manifest records the other versions, source hashes and timings. This experiment was executed in the assistant workspace on CPU. A matching notebook is provided for Colab; an actual user Colab session and Google Slides import still need confirmation.

## 6. Results

Validation metrics below explain the selection. Threshold values were chosen on this sample, so threshold-dependent validation metrics are in-sample tuning diagnostics.

{val}

The primary test comparison uses raw probabilities and each model's validation-selected threshold:

{results}

Confusion counts and additional metrics:

{counts}

![Test ROC](../figures/roc_curves.png)

The selected {pretty(m['selected_model'])} reaches test AUC {selected.roc_auc:.4f}, accuracy {selected.accuracy:.4f} and recall {selected.recall:.4f}. The Dummy prior reaches accuracy {dummy.accuracy:.4f} and AUC 0.5000. The results do not establish a meaningful ranking improvement. There is no multi-seed or account-cluster uncertainty study, and we do not claim statistical significance from small AUC differences. More complex models do not justify a deployment claim in this experiment.

### 6.1 Calibration and decision thresholds

We fit a positive-slope sigmoid on the selected model's validation logit scores, retaining the exact frozen base model. Test Brier score changes from {c.loc['raw','brier']:.4f} to {c.loc['calibrated','brier']:.4f}; the prior control scores {dummy.brier:.4f} (lower is better). Calibration can improve probability scale while leaving ordering unchanged, and it does not demonstrate stronger discrimination. We do not refit the classifier after calibration. The raw threshold {m['raw_decision_threshold']:.2f} maps to calibrated threshold {m['calibrated_decision_threshold']:.4f}, preserving the classification rule; calibration is not a retuning of the test decisions.

![Reliability](../figures/reliability.png)

### 6.2 Outcome-leakage control

A depth-one tree trained on close_value is intentionally invalid for pre-close prediction. Lost records have zero close value, making this post-outcome field an answer key in these files. leakage_audit.csv compares its test performance with the selected honest model on the same split. This control is never used for model selection or open-deal scoring. It explains why apparently excellent scores from outcome columns must be rejected.

![Leakage control](../figures/leakage_audit.png)

## 7. Explanations and decision-support outputs

We export native LR coefficients, RF impurity importance and TabNet attention importance with their distinct meanings. Held-out permutation importance measures the AUC change when one original feature is shuffled. Correlated inputs can share information, so low or negative permutation importance is possible and is not evidence that a feature can never matter.

For each open deal, the selected calibrated model reports reference sensitivity: delta_j = P(Won | observed inputs) minus P(Won | input j replaced by its training reference). References are training medians for numeric fields and training modes for categories. The three largest positive and negative changes are shown; all deltas are saved in JSON. These independent substitutions are not additive SHAP values and are not causal effects. Correlated fields can create unrealistic combinations when only one is replaced. The display explains model behavior, not what a sales representative should change.

Separately, TreeSHAP [6] explains the raw Random Forest on {shap['rows']} dated test rows. The saved expected probability plus per-feature SHAP values reconstructs the raw RF prediction with maximum error {shap['max_additivity_error']:.2e}. This provides an additive diagnostic for that RF only; it is not relabeled as an explanation of the selected calibrated classifier.

![RF explanations](../figures/rf_shap.png)

The output contains {m['scored_open_rows']:,} rows with ID, context, P(Won), P(Lost), predicted outcome, priority, input warnings, model name and explanation method. {example_text} This is a snapshot demonstration; the model may have been fitted using events later than an open record's original engagement date. It is not a claim that a probability was available on that historic date.

Priority labels use fixed illustrative cutoffs: Low below 0.40, Medium from 0.40 to below 0.70, High at least 0.70. These are separate from the learned class threshold. Test group diagnostics are:

{priority.to_markdown(index=False,floatfmt='.4f')}

Small or empty groups provide little evidence, and observed rates need not increase across bands. We include validation sensitivity tables for alternative cutoffs, but do not claim a business benefit, expected revenue or intervention uplift. Better account coverage and prospectively collected sales activities are necessary before a real prioritization policy can be evaluated.

## 8. Limitations, completion and next work

{limitations}

The promised four classifiers, full metric set, EDA/cleaning, feature importance and optional SHAP have executable implementations and generated evidence. The repository includes a Colab entry notebook, readable instructions, editable paper and slide drafts, and an ESL presentation script for three speakers. The software work is complete for this experimental scope; completion does not mean that the predictive model is useful in production. Slides satisfy the documented slides-and/or-video deliverable; no video recording is claimed. No autonomous agent or production service was required in the CRM proposal.

The next external steps are for the group to run the notebook in its own Colab account, review slides after Google Slides import, verify member contributions, rehearse and submit according to the professor's schedule. The code is prepared on a feature branch for peer review through a pull request to the group repository. Submission and merge are separate team actions. Research extensions include multiple temporal windows, account-group evaluation, more informative pre-close activity features, separate calibration data and prospective measurement of a specified sales policy. Such work is future work, not a reported completed experiment.

## Reproducibility

Run `python scripts/setup_cpu.py`, then `.venv-crm/bin/python -m src.run_project`. The default output is reports/crm/final. Use `.venv-crm/bin/python -m pytest -q` for tests. The upload-first notebook invokes the same entry point. `--quick` runs small neural budgets and writes SMOKE_TEST_NOT_FINAL artifacts separately. Each run saves source/data hashes, split assignments, configurations, raw test probabilities and a serialized frozen scoring model. Only load the model bundle generated by this trusted project; joblib is not an untrusted-file format.

## References

'''+'\n\n'.join(REFERENCES)+'\n'


def markdown_to_docx(text,outdir):
    doc=Document(); section=doc.sections[0]
    section.top_margin=DInches(.7); section.bottom_margin=DInches(.7)
    doc.styles['Normal'].font.name='Arial'; doc.styles['Normal'].font.size=DPt(10)
    doc.styles['Normal'].paragraph_format.space_after=DPt(6)
    lines=text.splitlines(); i=0
    while i<len(lines):
        line=lines[i].strip(); i+=1
        if not line: continue
        if line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):
                rows.append(lines[i].strip()); i+=1
            parsed=[[v.strip() for v in row.strip('|').split('|')] for row in rows]
            parsed=[r for r in parsed if not all(set(v)<=set(':- ') for v in r)]
            table=doc.add_table(rows=0,cols=len(parsed[0])); table.style='Light Shading Accent 1'
            for row in parsed:
                for cell,value in zip(table.add_row().cells,row):
                    cell.text=value
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs: run.font.size=DPt(8)
        elif line.startswith('!['):
            rel=line.split('](',1)[1].rstrip(')')
            doc.add_picture(str(outdir/rel),width=DInches(5.7))
            doc.add_paragraph(line[2:].split(']',1)[0],style='Caption')
        elif line.startswith('#'):
            level=len(line)-len(line.lstrip('#'))
            doc.add_heading(line.lstrip('# ').replace('**',''),level=min(level-1,2))
        else:
            doc.add_paragraph(line.removeprefix('- ').replace('**','').replace('`',''),style='List Bullet' if line.startswith('- ') else None)
    doc.save(outdir/'CRM_Final_Report.docx')


def slide_content(m,q,t,c,example):
    s=m['split']; selected=t.set_index('model').loc[m['selected_model']]
    cr=c.set_index('variant'); name=pretty(m['selected_model'])
    if example is None:
        example_lines=['No eligible Engaging opportunities in this run','1. Probability output: empty CSV with its schema','2. Input-factor output: no eligible records to explain','No invented example, score or outcome']
        example_note='The system normally provides win and loss probabilities together with the main input factors. In this run, there are no eligible Engaging opportunities to score. The exported CSV still has its columns, but it has no data rows. We do not invent an example or assign a result to a Prospecting record. The model comparison and test evaluation are still available from the closed opportunities. This makes the system behavior clear when there is no open work to score.'
    else:
        example_lines=[f"Actual open record: {example.opportunity_id} / {example['product']}",f'1. Win {example.win_probability:.1%}  |  Loss {example.loss_probability:.1%}',
               '2. Main positive and negative input factors',f'Predicted class: {example.predicted_outcome}; priority: {example.priority}',
               'Outcome unknown • frozen-model snapshot demonstration']
        example_note=f'The first output is the estimated win probability and its complementary loss probability. This real open record has about {example.win_probability:.0%} predicted win probability. The second output explains the main input factors affecting that score. The CSV also includes a predicted class, a simple priority label and missing-data warnings. We should be careful with this example: the actual outcome is unknown, and this is a snapshot demonstration. We are not claiming that this score was available on the original engagement date, or that the priority label has proven business value.'
    return [
      dict(title='CRM sales opportunities',kicker='CS 582 / GROUP 2',
        lines=['Predicting Won vs Lost — and explaining the score','An auditable CPU workflow for sales decision support','Main finding: the available inputs provide weak predictive signal.'],
        note='Hello everyone. Our project is about predicting whether a CRM sales opportunity will be won or lost. We also want to show which inputs affect the model score. We have now completed the comparison of the models in our proposal. The main finding is that the available data gives us only a weak prediction signal. So today we will explain both the working system and its limits. We will start with the data, then show the evaluation, and finish with the outputs and next steps.'),
      dict(title='What goes into the decision?',kicker='01 / INPUTS',
        lines=['Opportunity: product, sales agent and engagement date','Account: sector, revenue, employees, location and age','Product + team: list price, series, manager and region','18 predictors; outcome fields and opportunity ID are excluded.'],
        note='We use four connected tables. The opportunity table gives us the product, sales agent and engagement date. The account table adds information such as industry, revenue and company size. The other tables add product price and sales team information. After preparation, we use eighteen predictors. We do not use the final close value, close date or deal stage as prediction inputs, because they can reveal the answer. We keep the opportunity ID only to connect a prediction back to its original record.'),
      dict(title='Data prepared and checked',kicker='02 / DATA QUALITY',
        lines=[f"{q['raw_rows']['sales_pipeline']:,} opportunities → {q['rows']:,} observed Won/Lost outcomes",f"{q['scorable_open_rows']:,} Engaging records scored; {q['open_rows']-q['scorable_open_rows']:,} Prospecting excluded",'Fixed product names; validated keys, dates and joins',f"{q['scorable_missing_account']:,} scored records have missing account information."],figure='eda_missingness.png',
        note=f"The data contains {q['raw_rows']['sales_pipeline']:,} opportunities. We use the {q['rows']:,} closed records for supervised learning. The remaining records are still open. We score the {q['scorable_open_rows']:,} Engaging records, but leave out the {q['open_rows']-q['scorable_open_rows']:,} Prospecting records. We also fixed inconsistent product names and checked the table joins. An important issue is missing account information: {q['scorable_missing_account']:,} scored records have this warning. We make that visible because missing inputs can reduce how much we can trust a model score."),
      dict(title='Respect what was known at the time',kicker='03 / EVALUATION DESIGN',
        lines=[f"Train: {s['train_rows']:,} rows; labels known before {s['validation_start']}",f"Validation: {s['validation_rows']:,} rows; labels known before {s['test_start']}",f"Test: {s['test_rows']:,} later engagements; {s['purged_rows']:,} late labels purged",'Same-date deals stay together; this is an exploratory holdout.'],
        note='We split by engagement date, but date order alone is not enough. Imagine a deal starts in June but closes in October. Its outcome was not known in July, so it should not train a model used at that time. We exclude records like this from the earlier training or validation period. Every model uses the same split. We also keep deals from the same date together. This is still an exploratory historical test, because we have already studied this dataset. Now my teammate will explain the models and results.'),
      dict(title='Compare all promised models',kicker='04 / METHODS',
        lines=['Control: Dummy prior predicts the training win fraction','Baselines: Logistic Regression and Random Forest','Advanced models: MLP (64, 32) and TabNet','Train-only preprocessing; validation selects model and threshold.'],
        note='We trained all four models from our proposal, plus a simple control model. The control always gives the training win rate. Logistic Regression is our linear baseline, and Random Forest can learn more complex patterns. We also trained an MLP and TabNet on CPU. All preprocessing is fitted on the training set only. Validation data controls early stopping, model selection and the decision threshold. We did not choose the model using the test results. The saved settings let us repeat the comparison.'),
      dict(title='Results: weak discrimination',kicker='05 / MEASURED RESULTS',
        table=t[['model','accuracy','f1','roc_auc','brier']],
        lines=[f'{name} selected on validation; test AUC = {selected.roc_auc:.4f}','A majority-Won control can have high accuracy/F1 without useful ranking.'],
        note=f'Here are the actual test results. {name} was selected on validation, and its test AUC is about {selected.roc_auc:.2f}. An AUC of point five is the no-ranking baseline, so this result is only slightly above that level. The more advanced models did not give a convincing improvement in this run. Also, a model that predicts mostly Won can have a fairly high accuracy or F1 because Won is the larger class. That is why we show the control and several metrics together.'),
      dict(title='A perfect score can be the wrong answer',kicker='06 / LEAKAGE AUDIT',
        lines=['close_value is known after the outcome','A one-split tree nearly reveals Won/Lost from this field','The leaky control is excluded from all real scoring','Lesson: feature availability matters as much as model choice.'],figure='leakage_audit.png',
        note='We also ran a deliberately incorrect experiment to show the effect of data leakage. In these files, lost deals have zero close value. If we use that field, even a very small decision tree can almost directly recover the outcome. The score looks excellent, but the information is only known after the deal closes. So this would not help us predict a new opportunity. We keep this experiment only as an audit example. It is never used in the real prediction output.'),
      dict(title='Show what influences the score',kicker='07 / EXPLANATIONS',
        lines=['Selected model: compare each input with its training reference','Show the largest positive and negative probability changes','Separate RF TreeSHAP audit verifies additive explanations','Model associations are not causes or guaranteed sales actions.'],figure='rf_shap.png',
        note='For each scored opportunity, we show which inputs increase or decrease the model score compared with a training reference. For example, we replace one numeric value with its training median and measure the probability change. These changes help explain the model, but they are not additive SHAP values. Separately, we computed TreeSHAP for the raw Random Forest and checked that its contributions reconstruct the prediction. We clearly label the two methods. Neither method proves cause and effect. My teammate will now show the output and practical limits.'),
      dict(title='Two outputs, with an honest context',kicker='08 / OUTPUT EXAMPLE',
        lines=example_lines,note=example_note),
      dict(title='Calibration helps scale, not ranking',kicker='09 / LIMITS',
        lines=[f"Test Brier: {cr.loc['raw','brier']:.4f} raw → {cr.loc['calibrated','brier']:.4f} calibrated",'Calibration uses validation; the base model stays frozen','Priority cutoffs 0.40 / 0.70 are illustrative and need validation','Closed-only sample, time drift and missing accounts remain risks.'],figure='reliability.png',
        note='We used validation data to adjust the probability scale. On the test set, we compare the Brier score before and after calibration. A lower score is better, but calibration does not improve the ordering of deals. The base model stays frozen so the calibration still matches it. Our priority cutoffs are only examples and need business validation. We also have limits from using only observed closed outcomes, changes over time and missing account information. So this version should be treated as a research demonstration, not a production sales tool.'),
      dict(title='What is complete?',kicker='10 / DELIVERY STATUS',
        lines=['Data checks, EDA and availability-aware split','LR, RF, MLP, TabNet + control; full metrics and confusion matrices','Local explanations, feature importance, SHAP and leakage audit','One-command run, notebook, README, paper and editable slides'],
        note='The main items from our proposal now have working code and generated outputs. We have data checks, exploratory analysis, all four promised models, the comparison metrics and confusion matrices. We also completed feature importance and the optional SHAP analysis. The workflow creates the open-deal CSV, figures, report and slides from the same run. This helps keep our presentation consistent with the code. The remaining checks involve the group account and final presentation review. We still need to confirm the notebook in our own Colab session and check the slides after import.'),
      dict(title='Demo, review and next steps',kicker='11 / HANDOFF',
        lines=['Demo: open notebook → Run all → inspect results → download ZIP','Team review: paper, slide layout, speaker roles and limitations','Code review: review the feature-branch PR before merging','Future research: better pre-close activity data and prospective tests'],
        note='For the demo, we can upload the project ZIP, run the notebook and inspect the generated results. We should check the run status, compare the models and open one scored opportunity with its explanation. Before submission, each member should review the report, rehearse their part and confirm the instructor schedule. The code is prepared for review in a pull request to our group repository. We should review the changes before merging. For future research, more useful pre-close activity data and a prospective test would matter more than simply adding another complex model. Thank you. We are happy to take questions.'),
    ]


def build_slides(slides,out,m):
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    ink='122A3A'; teal='127D88'; muted='506574'; bg='F4F7F8'
    def rect(sl,x,y,w,h,color):
        shape=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
        shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor.from_string(color); shape.line.fill.background()
        return shape
    def text(sl,value,x,y,w,h,size=22,color=ink,bold=False):
        box=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame
        tf.word_wrap=True; tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
        for i,line in enumerate(value.split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line
            p.font.name='Arial'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color)
            p.space_after=Pt(12)
        return box
    for i,item in enumerate(slides):
        sl=prs.slides.add_slide(prs.slide_layouts[6]); rect(sl,0,0,13.333,7.5,bg); rect(sl,0,0,.18,7.5,teal)
        text(sl,item['kicker'],.65,.38,12,.3,12,teal,True)
        text(sl,item['title'],.65,.96,12,1.0,32,ink,True)
        if 'table' in item:
            data=item['table']; table=sl.shapes.add_table(len(data)+1,5,Inches(.7),Inches(2),Inches(11.9),Inches(2.9)).table
            widths=[3.8,2,2,2,2.1]
            for col,width in zip(table.columns,widths): col.width=Inches(width)
            rows=[['Model','Accuracy','F1 (Won)','ROC-AUC','Brier ↓']]+[[pretty(row.model),f'{row.accuracy:.4f}',f'{row.f1:.4f}',f'{row.roc_auc:.4f}',f'{row.brier:.4f}'] for row in data.itertuples()]
            for r,row in enumerate(rows):
                for j,value in enumerate(row):
                    cell=table.cell(r,j); cell.text=value
                    cell.fill.solid(); cell.fill.fore_color.rgb=RGBColor.from_string(teal if r==0 else ('E4EEF0' if r%2 else 'FFFFFF'))
                    for p in cell.text_frame.paragraphs:
                        p.font.name='Arial'; p.font.size=Pt(18); p.font.bold=r==0; p.font.color.rgb=RGBColor.from_string('FFFFFF' if r==0 else ink)
            text(sl,'\n'.join(item['lines']),.75,5.25,11.9,1.1,20)
        elif 'figure' in item:
            for j,line in enumerate(item['lines']):
                rect(sl,.7,2.07+j*1.03,.07,.42,teal)
                text(sl,line,.92,2.02+j*1.03,4.9,.9,21)
            from PIL import Image
            path=out.parent/'figures'/item['figure']
            with Image.open(path) as img: width,height=img.size
            w=min(6.6,4.55*width/height); h=w*height/width
            sl.shapes.add_picture(str(path),Inches(6.05+(6.6-w)/2),Inches(2.0+(4.55-h)/2),width=Inches(w),height=Inches(h))
        else:
            for j,line in enumerate(item['lines']):
                rect(sl,.7,2.15+j*.88,.08,.36,teal)
                text(sl,line,.98,2.08+j*.88,11.5,.8,26 if i==0 else 24)
        footer=f"CS 582 • Group 2 | {m['mode']} | {i+1:02d}/12"
        text(sl,footer,.65,7.05,12,.25,10,muted)
        sl.notes_slide.notes_text_frame.text=f"Suggested speaker: {TEAM[i//4]}\n\n{item['note']}\n\nEvidence: reports/crm/final; run {m['started_utc']}. Sources: Maven dataset [1]; methods [4–7], references in the paper."
    prs.save(out/'CRM_Final_Presentation.pptx')


def build_deliverables(output_dir):
    root=Path(output_dir); out=root/'deliverables'; out.mkdir(exist_ok=True)
    m=json.loads((root/'run_manifest.json').read_text()); q=json.loads((root/'data_quality.json').read_text())
    t=pd.read_csv(root/'test_metrics.csv'); v=pd.read_csv(root/'validation_metrics.csv'); c=pd.read_csv(root/'calibration_test.csv')
    priority=pd.read_csv(root/'priority_test.csv'); history=json.loads((root/'model_training.json').read_text())
    shap=json.loads((root/'shap_audit.json').read_text()); scored=pd.read_csv(root/'open_deal_predictions.csv')
    if len(scored)!=m['scored_open_rows'] or len(scored)!=q['scorable_open_rows']:
        raise ValueError('Score count does not match run manifest/data-quality report')
    missing=int(scored.account_missing.sum())
    if missing!=q['scorable_missing_account']:
        raise ValueError('Missing-account count does not match the exported scores')
    complete=scored.loc[~scored.account_missing] if len(scored) else scored
    example=(complete if len(complete) else scored).iloc[0] if len(scored) else None
    paper=paper_text(m,q,t,v,c,priority,history,example,shap)
    (out/'CRM_Final_Report.md').write_text(paper,encoding='utf-8'); markdown_to_docx(paper,out)
    slides=slide_content(m,q,t,c,example); build_slides(slides,out,m)
    script='# Natural ESL presentation script\n\nSuggested allocation: 12 slides / 3 members / 4 slides each. About 8–10 minutes total at a comfortable speaking pace. Swap the names if the group prefers; these are speaking assignments, not claims about past contributions.\n\n'
    for i,item in enumerate(slides):
        script+=f"## Slide {i+1}: {item['title']}\n\n**{TEAM[i//4]}**\n\n{item['note']}\n\n"
    script+='''## Short Q&A practice

**What is the novelty?** Our contribution is an explainable CRM workflow with checks for future information and missing inputs. It combines probability prediction with reasons and an audit trail. We are not claiming a new machine learning algorithm.

**Why is the accuracy low?** The available pre-close features contain little useful signal in this test. We prefer to report that honestly instead of using information that is only known after the outcome.

**Why did you use TabNet?** We promised to compare traditional models with an advanced tabular model. TabNet gives us that comparison, but it did not show a useful improvement here.

**Can the sales team use the priority score now?** It is a demonstration at this stage. We need better data coverage and a prospective test before using it to allocate sales effort.

**What is the difference between SHAP and your local explanation?** Our main output changes one input to a training reference and measures the score change. Those changes do not add up to the prediction. TreeSHAP is a separate additive explanation for the raw Random Forest.

**What have you actually run?** We ran all five models on CPU in the development environment. The notebook uses the same code. The group still needs to confirm the run in its own Colab account.

**Is this an intelligent agent?** This version is a supervised decision-support pipeline. The proposal did not require an autonomous agent, and we do not claim one.

**Does a 70 percent score mean this particular deal will win?** No. It is a model estimate, not a guarantee. Calibration, missing inputs and changes in the data all affect how much we can trust it.
'''
    (out/'SPEAKER_SCRIPT_ESL.md').write_text(script,encoding='utf-8')
    scriptdoc=Document(); scriptdoc.add_heading('CS 582 — Speaking notes',0)
    for line in script.splitlines():
        if line.strip(): scriptdoc.add_paragraph(line.replace('**','').lstrip('# '))
    scriptdoc.save(out/'SPEAKER_SCRIPT_ESL.docx')
    summary=f"# Results from the current run\n\nMode: **{m['mode']}**. Selected on validation: **{pretty(m['selected_model'])}**.\n\n"+t[['model','accuracy','f1','roc_auc','brier']].to_markdown(index=False,floatfmt='.4f')
    summary+=f'\n\nAll five models ran. Open scores: {len(scored):,}; missing-account warnings: {missing:,}. Near-chance AUC does not support production prioritization. Calibration and priority bands are diagnostics, not proven business improvement. See the paper for limitations.\n'
    (out/'RESULTS_SUMMARY.md').write_text(summary,encoding='utf-8')
    (out/'TEAM_REVIEW.md').write_text('''# Final team review before submission

- [ ] Run the upload-first notebook in the group Colab account; save the results ZIP.
- [ ] Open the PPTX in Google Slides and inspect fonts, charts and speaker notes.
- [ ] Each member reads and can explain the data exclusions, split and actual results.
- [ ] Verify names and agree the real member contribution statement; none is fabricated here.
- [ ] Check the professor's exact presentation slot, paper format and submission channel.
- [ ] Rehearse the 12-slide final talk (4 slides per member). The earlier 1–8-slide limit was for the progress forum; confirm any final-talk limit separately.
- [ ] Review the feature-branch pull request against the group repository before merging.
- [ ] Submit paper, slides and code. No video recording or submission has been performed by this script.
''',encoding='utf-8')
    print(f'Generated paper, slides and speaker notes: {out}',flush=True)


if __name__=='__main__':
    parser=argparse.ArgumentParser(); parser.add_argument('output_dir',type=Path)
    build_deliverables(parser.parse_args().output_dir)
