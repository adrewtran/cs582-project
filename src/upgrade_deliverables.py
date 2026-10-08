"""A/B evidence, paper, editable presentation and timed natural ESL narration.

The slide template is authored with Artifact Tool. A standard-library OOXML
fill keeps regeneration free and self-contained on Colab (no design API).
"""
import argparse
import json
from pathlib import Path
import re
import zipfile
from xml.sax.saxutils import escape
import pandas as pd
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from src.deliverables import TEAM,REFERENCES,pretty
from src.presentation_story import story,result_narrative

ROOT=Path(__file__).resolve().parents[1]


def table(frame,columns):
    def display(v):
        return f'{v:.4f}' if isinstance(v,float) else str(v)
    rows=[columns,['---']*len(columns)]+[[display(v) for v in row] for row in frame[columns].itertuples(index=False,name=None)]
    return '\n'.join('| '+' | '.join(row)+' |' for row in rows)


def report(m,q,ablation,example,calibration,coverage):
    selected=ablation.loc[(ablation.experiment==m['selected_experiment']) & (ablation.model==m['selected_model'])].iloc[0]
    raw=ablation.loc[(ablation.experiment=='raw') & (ablation.model==m['selected_model'])].iloc[0]
    interpretation=result_narrative(m,ablation)
    s=m['split']; available=example.get('available',True)
    if available:
        p=example['model_prediction']['win_probability']; opp=example['opportunity']
        demo_description=(f"The default example is opportunity {opp['opportunity_id']}, account {opp['account']}, product {opp['product']}. "
            f"The saved model estimates win probability {p:.2%} and loss probability {1-p:.2%}, "
            f"with loss risk {example['agent_recommendation']['loss_risk']}. Its real outcome remains unknown. "
            f"Scoring context: {example['timing']['scoring_context']}. The example is selected by date and ID, not a high score.")
    else:
        demo_description='No dated Engaging opportunity is available. No probability or recommendation is invented; the empty export is explicit.' 
    parts=[f'''# CRM Sales Decision Support

CS 582 Group 2

{', '.join(TEAM)}

## Abstract

We evaluate whether strictly prior sales history improves CRM win prediction and connect predictions to an inspectable sales assistant. We compare six CPU classifiers on the same availability-aware split, using 18 raw features (A) or 37 raw and historical features (B). Validation selects {m['selected_experiment']} / {pretty(m['selected_model'])}. Its exploratory test ROC-AUC is {selected.test_roc_auc:.4f}, compared with {raw.test_roc_auc:.4f} for the same model using raw inputs. {interpretation['conclusion']} Our contribution is the combination of a leakage audit, a tested temporal feature engine and an evidence-based rule agent. All metrics come from the {m['mode']} run beginning {m['started_utc']}.

## 1 Problem and related work

The model estimates P(Won) and its complement among opportunities whose eventual outcome is observed. Open deals are never assigned artificial Lost labels. We do not predict a fixed closing horizon. The intended use is to inspect a score and decide what information needs review, with a person retaining control of sales actions.

Sales win propensity prediction and probability-based B2B workflows already exist [2,3]. TabNet provides an attention-based tabular comparator [4,7]. CatBoost provides a categorical gradient-boosting comparator [8]. We apply existing algorithms rather than claim a novel model architecture. Explainability supports inspection, but SHAP alone is not our novelty.

## 2 Data and preprocessing

The supplied Maven CRM dataset describes a fictitious B2B hardware company [1]. It contains {q['raw_rows']['sales_pipeline']:,} opportunities, {q['raw_rows']['accounts']} accounts, {q['raw_rows']['products']} products and {q['raw_rows']['sales_teams']} team records. A fifth CSV is the data dictionary. There are {q['rows']:,} closed records with {q['positives']:,} Won and {q['negatives']:,} Lost outcomes. The system scores {q['scorable_open_rows']:,} dated Engaging opportunities. It excludes undated Prospecting records. Among scored opportunities, {q['scorable_missing_account']:,} lack account information.

We validate key uniqueness, many-to-one joins, stage values, nonnegative numeric fields and valid date ordering. We normalize GTXPro to GTX Pro and the sector spelling technolgy to technology. The source CSV files remain unchanged. run_manifest.json records their SHA-256 hashes and the source hashes for this run.

The 18 raw inputs include product and team information, account attributes, list price and engagement calendar fields. Account revenue is in millions in the source. Revenue per employee converts it to currency units first. Current close_value, close_date, deal_stage and opportunity_id are never model predictors. Stage defines the target; dates enforce availability; IDs link audit records. Histories use other deals' already-observed outcomes, not the current outcome.

One-hot categorical encoders and median imputation fit only training rows. LR, MLP and TabNet scale numerics using training statistics. CatBoost uses native categorical strings and training-fitted numeric medians. Missing categories remain explicit. Static account, product and team snapshots may not describe each entity's true historical state.

## 3 Temporal design and feature engineering

Training has {s['train_rows']:,} rows, validation {s['validation_rows']:,}, and test {s['test_rows']:,}. We purge {s['purged_rows']:,} records whose outcome was not available before the next period. Validation starts {s['validation_start']} and test starts {s['test_start']}. All six models and both feature variants receive the same row IDs. Win fractions are {s['train_win_rate']:.2%}, {s['validation_win_rate']:.2%} and {s['test_win_rate']:.2%}, respectively.

The history engine uses a frozen training archive. At engagement time T, a record can contribute only if close_date < T and its opportunity ID differs from the query. Same-day closures are excluded because intraday order is unknown. Validation and test outcomes never enter the archive, even when they have already closed. This conservative design is an offline comparison, not an online adaptive evaluation.

For each sales agent, account and product, we derive a prior closed count, smoothed win rate, mean observed close value, mean cycle length and cold-start flag. Four global prior statistics complete the 19 added features. The smoothed entity win rate is (prior entity wins + 5 times eligible global win rate) divided by (prior entity count + 5). The smoothing strength is fixed before testing. Missing entity keys use the eligible global prior with local count zero. With no eligible global history, the probability prior is 0.5, count/value/cycle defaults are zero, and cold-start flags remain visible.

The row-level audit records query date, latest eligible closure and archive size. Unit tests check same-day and self exclusion, shuffled input order, future-outcome invariance and held-out-label invariance. The following coverage table counts queries with at least one eligible entity-specific past closure.

{table(coverage,['split','entity','rows','with_history','coverage'])}

## 4 Model selection and experimental protocol

A uses raw inputs. B adds histories. We fit Dummy prior, Logistic Regression, Random Forest, MLP, TabNet and CatBoost for both variants. Dummy returns the training win fraction. LR uses L2 regularization and C=1. RF uses 300 trees and minimum leaf size 5. MLP uses hidden layers 64 and 32, alpha 0.001, learning rate 0.001, at most 120 epochs and patience 15. TabNet uses n_d=n_a=8, three steps, learning rate 0.02, at most 80 epochs and patience 12. CatBoost uses depth 5, learning rate 0.05, up to 500 iterations, L2 leaf regularization 5 and validation early stopping 40. CPU threads are limited and the random seed is 42. Smoke runs use smaller budgets and are explicitly unreportable.

Validation ROC-AUC selects the model and experiment, with lower validation Brier score and deterministic name ordering breaking ties. We retain validation-selected neural checkpoints. A validation macro-F1 search chooses class thresholds. Positive-slope sigmoid calibration uses validation and keeps the base classifier frozen. selection_lock.json saves both experiments' choices before either evaluates test predictions. The validation sample serves multiple roles, so these estimates can be optimistic. Test results do not trigger another configuration search.

Diagnostic C is a separate depth-one tree using the current final close value. It illustrates invalid outcome leakage and never supplies real opportunity scores. We report AP as average precision, rather than calling it trapezoidal PR-AUC. Precision, recall and F1 treat Won as positive. Confusion counts and additional metrics appear in the appendix and CSV files.

## 5 Validation and test results

Validation comparisons determine selection. These numbers are tuning diagnostics rather than an independent estimate of generalization.

{table(ablation,['experiment','model','validation_roc_auc','validation_brier','validation_threshold'])}

The test comparison below uses the original probability scale of each model and its validation-selected threshold. Lower Brier and log loss are better. AUC 0.5 represents no useful ranking.

{table(ablation,['experiment','model','test_accuracy','test_f1','test_roc_auc','test_brier'])}

The selected model changes from raw-feature AUC {raw.test_roc_auc:.4f} to selected-variant AUC {selected.test_roc_auc:.4f}, a descriptive difference of {selected.test_roc_auc-raw.test_roc_auc:+.4f}. {interpretation['conclusion']} The history representation may capture time trends and sparsity rather than stable opportunity information. Training history grows over time, while validation/test use a frozen archive; this changes the feature distribution. Those are plausible explanations, not separately verified causes. We did not run a multi-seed or cluster-bootstrap significance study.

The highest test AUC is a descriptive observation only. We retain the validation winner rather than choose a model after inspecting test results. Dummy can attain a high Won-class F1 by predicting the majority class, so F1 alone is not enough to demonstrate useful prioritization.

### Probability calibration

{table(calibration,['variant','roc_auc','brier','log_loss','threshold'])}

Calibration changes the probability scale while preserving ranking, apart from numerical ties. Its result does not establish reliable future probabilities. Brier must be interpreted against the prior-only baseline as well as the uncalibrated selected model.

## 6 Explanations and sales assistant

The selected calibrated model produces the win probability and complementary loss probability. Its local explanations replace one feature with a training median or mode and measure the probability difference. These are nonadditive reference sensitivities, not causal effects. Correlated history features can yield unrealistic single-feature replacements. RF TreeSHAP is a separate supporting diagnostic for the uncalibrated forest, with numerical reconstruction checks [6]. It is not mislabeled as an explanation of another selected classifier.

The deterministic assistant receives the probability, historical context, signed sensitivities and warnings. It returns loss risk HIGH below P(win)=0.40, MEDIUM from 0.40 to below 0.70, and LOW at 0.70 or above. Risk describes estimated loss probability; it is distinct from the legacy High win-priority label. Fixed bands are illustrative, not a validated allocation policy.

Rules request account completion when identity is missing, direct verification when fewer than five account outcomes exist, product-fit review when at least five past product outcomes lag the eligible global win rate by over 0.10, or a second sales review when high loss risk and sufficiently supported agent history agree. The output includes at most two triggered rules plus two universal checks: verify current opportunity status and review the evidence before action. Each action exposes its rule ID, observed evidence and rationale. No rule sends a message, changes a price, measures uplift, or claims an intervention will cause a win. No paid LLM API is required.

Our course-level contribution has three linked parts: an explicit leakage audit, an availability-tested historical feature engine and an inspectable rule-based sales assistant. CatBoost is a stronger comparator and SHAP is supporting analysis. Neither is claimed as our own new algorithm. The agent is decision support, not autonomous planning or a learned sales policy.

## 7 Live demonstration

{demo_description}

The CLI reloads model_bundle.joblib, reconstructs histories from its saved archive, recalculates the probability and reasons, and runs the rule engine. Tests compare it with the exported score. Records whose engagement predates model availability are explicitly marked retrospective snapshots. Even later records remain historical demonstrations, not a live CRM deployment. Only load joblib bundles from trusted sources.

## 8 Limitations and next evidence

The data has been inspected in earlier project iterations, so this is an exploratory holdout. Closed-only sampling and end-of-snapshot censoring remain. Accounts and agents recur across periods, so the experiment is not a test of entirely new customers. Input tables are static. The training archive remains frozen and early training examples have thin or no history. Open records with missing accounts differ from the complete-account labeled sample.

The measured discrimination is weak. Recommendations have logical rule tests but no user study or intervention-outcome evaluation. The rule thresholds and probability bands need prospective validation. We do not claim increased revenue, expected deal value, causal actions or production readiness. Future work should collect dated interaction, qualification and competitor information; evaluate rolling periods or an external cohort; reserve separate calibration data; and test whether users find the evidence-linked actions useful. Changes should be predeclared before consulting a new test.

## 9 Reproducibility and conclusion

Run python scripts/setup_cpu.py with Python 3.12, then .venv-crm/bin/python -m src.run_project. The CPU pipeline writes both experiments, the selection lock, all metrics, audits, models, 1-row live demo example and batch assistant outputs. The notebook invokes the same code. A successful software run is distinct from a useful predictive model. Our verified execution environment is development CPU; actual group Colab execution, Google Slides import and course submission require team confirmation.

The extension strengthens the experimental audit and connects model output to transparent review actions. It does not justify a claim that historical features improve predictive accuracy on this dataset. Reporting this distinction is central to the project.

## Appendix Complete metrics and confusion counts
''']
    for name in ['raw','history']:
        subset=ablation.loc[ablation.experiment.eq(name)].copy()
        subset['model']=subset.model.map(pretty)
        for columns in [
            ['model','test_precision','test_recall','test_macro_f1','test_balanced_accuracy'],
            ['model','test_average_precision','test_log_loss','test_tn','test_fp','test_fn','test_tp']]:
            clean=subset[columns].rename(columns={c:c.removeprefix('test_') for c in columns})
            parts.append(f'\n### {name.capitalize()} features\n\n'+table(clean,list(clean)))
    parts.append('\n## References\n\n'+'\n\n'.join(REFERENCES+[
        '[8] L. Prokhorenkova et al. CatBoost: unbiased boosting with categorical features. NeurIPS, 2018. https://arxiv.org/abs/1706.09516 ; implementation https://catboost.ai/docs/ (version 1.2.10 used).']))
    return '\n'.join(parts)


def write_docx(markdown,path):
    """Readable Word report with native tables, deliberate widths and repeating headers."""
    doc=Document(); section=doc.sections[0]
    section.page_width=Inches(8.5); section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.7)
    section.left_margin=section.right_margin=Inches(.7)
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
        style=doc.styles[name]; style.font.name='Arial'; style.font.color.rgb=RGBColor(0,0,0)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
    doc.styles['Normal'].font.size=Pt(11); doc.styles['Normal'].paragraph_format.space_after=Pt(7)
    doc.styles['Normal'].paragraph_format.line_spacing=1.1
    lines=markdown.splitlines(); i=0
    while i<len(lines):
        line=lines[i].strip(); i+=1
        if not line: continue
        if line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].strip().startswith('|'): rows.append(lines[i].strip()); i+=1
            values=[[v.strip() for v in row.strip('|').split('|')] for row in rows]
            values=[row for row in values if not all(set(v)<=set(':- ') for v in row)]
            t=doc.add_table(rows=0,cols=len(values[0])); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
            widths=[1.55]+[(7.1-1.55)/(len(values[0])-1)]*(len(values[0])-1)
            if values[0][0]=='experiment':
                widths=[.80,1.58]+[(7.1-2.38)/(len(values[0])-2)]*(len(values[0])-2)
                values[0][0]='Input'
            for col,w in zip(t.columns,widths): col.width=Inches(w)
            for rownum,values_row in enumerate(values):
                row=t.add_row()
                if rownum==0:
                    repeat=OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(repeat)
                for colnum,(cell,value,w) in enumerate(zip(row.cells,values_row,widths)):
                    cell.width=Inches(w); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cell.text=value.replace('validation_','val ').replace('test_','').replace('_',' ')
                    tc=cell._tc.get_or_add_tcPr()
                    shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'DCE8EC' if rownum==0 else ('F4F7F8' if rownum%2 else 'FFFFFF')); tc.append(shade)
                    borders=OxmlElement('w:tcBorders')
                    for edge in ['top','left','bottom','right']:
                        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
                    tc.append(borders)
                    margins=OxmlElement('w:tcMar')
                    for edge in ['top','left','bottom','right']:
                        e=OxmlElement('w:'+edge); e.set(qn('w:w'),'75'); e.set(qn('w:type'),'dxa'); margins.append(e)
                    tc.append(margins)
                    for para in cell.paragraphs:
                        para.alignment=WD_ALIGN_PARAGRAPH.LEFT if colnum==0 or (colnum==1 and len(widths)>1 and widths[1]>1.4) else WD_ALIGN_PARAGRAPH.CENTER
                        para.paragraph_format.space_after=Pt(2); para.paragraph_format.line_spacing=1
                        for run in para.runs: run.font.size=Pt(9); run.bold=rownum==0
            doc.add_paragraph()
        elif line.startswith('#'):
            level=len(line)-len(line.lstrip('#')); title=line.lstrip('# ').replace('**','')
            doc.add_paragraph(title,style='Title' if level==1 else f'Heading {min(level-1,3)}')
        else:
            doc.add_paragraph(line.replace('**','').replace('`',''))
    doc.save(path)


def fill_deck(slides,target,mode):
    template=ROOT/'assets/CRM_History_Agent_Template.pptx'
    if not template.exists(): raise FileNotFoundError('Missing committed presentation template: '+str(template))
    values={}
    for i,item in enumerate(slides,1):
        values.update({f'title{i}':item['title'],f'kicker{i}':item['kicker'],
            f'footer{i}':f'CS 582 Group 2   {mode}   {i:02d}/12',
            f'note{i}':f"Suggested speaker: {TEAM[item['speaker_index']]} ({item['seconds']} seconds).\n\n{item['note']}\n\nEvidence: {item['source']}"})
        for j,line in enumerate(item['lines']): values[f'line{i}_{j}']=line
        for r,row in enumerate(item.get('table',[])):
            for c,value in enumerate(row): values[f'cell{i}_{r}_{c}']=str(value)
    with zipfile.ZipFile(template) as source,zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as dest:
        for info in source.infolist():
            data=source.read(info.filename)
            if info.filename.endswith('.xml'):
                xml=data.decode('utf-8')
                xml=re.sub(r'\[\[([A-Za-z0-9_]+)\]\]',lambda match:escape(values[match.group(1)]),xml)
                data=xml.encode('utf-8')
            dest.writestr(info,data)


def build_deliverables(output_dir,template_required=True):
    root=Path(output_dir); dest=root/'deliverables'; dest.mkdir(exist_ok=True)
    m=json.loads((root/'run_manifest.json').read_text()); q=json.loads((root/'data_quality.json').read_text())
    a=pd.read_csv(root/'ablation_comparison.csv'); c=pd.read_csv(root/'calibration_test.csv')
    example=json.loads((root/'demo_example.json').read_text()); coverage=pd.read_csv(root/'history_coverage.csv')
    scored=pd.read_csv(root/'open_deal_predictions.csv')
    if len(scored)!=m['scored_open_rows'] or len(a)!=12: raise ValueError('Deliverables require complete aligned experiment outputs')
    leakage=pd.read_csv(root/'leakage_audit.csv')
    leakage_auc=float(leakage.loc[leakage.model.eq('LEAKY_close_value_stump'),'roc_auc'].iloc[0])
    interpretation=result_narrative(m,a)
    slides=story(m,q,a,example,c,leakage_auc)
    (dest/'slide_content.json').write_text(json.dumps(slides,indent=2,ensure_ascii=False),encoding='utf-8')
    paper=report(m,q,a,example,c,coverage)
    (dest/'CRM_Final_Report.md').write_text(paper,encoding='utf-8'); write_docx(paper,dest/'CRM_Final_Report.docx')
    script='# CRM presentation script\n\nSuggested speaking allocation only. This does not claim each member performed these development tasks. Total 10 minutes including a short live command. Rehearse and adjust the pace without hiding limitations.\n\n'
    elapsed=0
    for i,item in enumerate(slides,1):
        end=elapsed+item['seconds']; stamp=lambda t:f'{t//60}:{t%60:02d}'
        script+=f"## Slide {i} {item['title']}\n\n{TEAM[item['speaker_index']]} | {stamp(elapsed)} to {stamp(end)}\n\n{item['note']}\n\n"
        if i==9: script+='Demo cue: run the prepared src.demo command. Show the probability, one reason, and the evidence beneath one action. Allow about 20 seconds for the screen walkthrough.\n\n'
        elapsed=end
    script+=f'''## Natural ESL answers for questions

What is new here? Our contribution combines a tested time-aware history engine, a leakage audit, and a sales assistant that explains each suggested review action. We use existing learning algorithms. We do not claim a new algorithm or proven sales improvement.

Did history improve accuracy? {interpretation['conclusion']} We report that result directly.

Why not select the highest test score? The highest observed test AUC is {interpretation['best_auc']:.4f} from {interpretation['best']}, but the selection rule uses validation. Changing the rule after seeing test results would make our evaluation less trustworthy.

Is this an agent or a chatbot? It is a deterministic decision-support agent. It observes a prediction and evidence, applies explicit rules, and recommends review actions. It has no LLM, autonomous planning, or ability to contact customers.

Why are past close values allowed? They belong to other training deals that closed before the query date. The current opportunity's final value is never used by the legitimate predictor. The audit records the latest eligible past closure.

What does the score mean? It is an estimated probability from a model trained on observed closed deals. It is not a guarantee for this opportunity, and it may not transfer to open deals with missing information.

Does the explanation prove why a sale is lost? No. It shows model sensitivity to input values. It does not prove a real causal effect or that changing the input will improve the outcome.

Why are the results weak? This snapshot has limited pre-close information. Historical sparsity, changing feature distributions and time drift may contribute. We have not isolated those causes experimentally.

How do you validate the agent? Tests check whether each rule fires under its stated evidence, whether actions are deterministic, and whether the live model matches the exported score. A user study and prospective intervention study remain future work.

Can we call this excellent? The implementation adds testable novelty and a reproducible demonstration. The professor decides the grade. Weak predictive performance and unvalidated business value remain limitations, so we should not promise an excellent grade.
'''
    (dest/'SPEAKER_SCRIPT_ESL.md').write_text(script,encoding='utf-8'); write_docx(script,dest/'SPEAKER_SCRIPT_ESL.docx')
    if template_required: fill_deck(slides,dest/'CRM_Final_Presentation.pptx',m['mode'])
    summary=f"# Current CRM results\n\nMode {m['mode']}. Validation selection: {m['selected_experiment']} / {m['selected_model']}.\n\n"
    summary+=table(a,['experiment','model','validation_roc_auc','test_roc_auc','test_accuracy','test_f1','test_brier'])
    summary+=f"\n\nOpen scores: {len(scored):,}; missing-account warnings: {int(scored.account_missing.sum()):,}. {interpretation['conclusion']} The agent is evidence-linked decision support, not a proven intervention.\n"
    (dest/'RESULTS_SUMMARY.md').write_text(summary,encoding='utf-8')
    (dest/'TEAM_REVIEW.md').write_text('''# Team checks before submission

- Run the notebook in the group Colab account and keep the results ZIP.
- Inspect all slides and notes after Google Slides import.
- Rehearse the 10-minute script and the real-model demo.
- Review the exact negative result and limitations together.
- Confirm actual member contributions and the professor's formatting, slide limit and presentation slot.
- Review the PR before merging. No merge, video recording or course submission occurs automatically.
''',encoding='utf-8')
    print('Generated upgraded report, presentation content and 10-minute script:',dest,flush=True)


if __name__=='__main__':
    parser=argparse.ArgumentParser(); parser.add_argument('output_dir',type=Path)
    parser.add_argument('--content-only',action='store_true')
    args=parser.parse_args(); build_deliverables(args.output_dir,not args.content_only)
